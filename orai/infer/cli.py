"""命令行：把文稿逐段改写，降低 AI 痕迹。

python -m orai.infer.cli 输入.txt|输入.docx -o 输出.txt|输出.docx --adapter outputs/dpo-r1/final \
    [--style 书面|口语] [--n 6] [--rounds 2] [--no-judge]

流程：按段切分 → 每段采样 n 个候选 → 本地裁判团（检测分 + 语义门槛）挑最优 →
仍像 AI 的段落再改一轮 → 写出结果和逐段报告（同名 .report.jsonl）。
裁判团和改写模型分时占用显存：先生成全部候选，释放后再打分。
"""

from __future__ import annotations

import argparse
import gc
import json
from pathlib import Path

MIN_LEN = 40  # 太短的段落（标题、图注）不改


def read_paragraphs(path: Path) -> list[str]:
    if path.suffix == ".docx":
        import docx
        return [p.text for p in docx.Document(str(path)).paragraphs]
    return path.read_text(encoding="utf-8").split("\n")


def write_output(src: Path, dst: Path, paras: list[str]):
    if dst.suffix == ".docx" and src.suffix == ".docx":
        import docx
        doc = docx.Document(str(src))
        for p, new in zip(doc.paragraphs, paras):
            if p.text != new and p.runs:
                p.runs[0].text = new  # 保留第一个 run 的格式
                for r in p.runs[1:]:
                    r.text = ""
        doc.save(str(dst))
    else:
        dst.write_text("\n".join(paras), encoding="utf-8")


def free_gpu():
    import torch
    gc.collect()
    torch.cuda.empty_cache()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("-o", "--output")
    ap.add_argument("--adapter", default="outputs/sft-v2/final")
    ap.add_argument("--style", default="书面", choices=["书面", "口语"])
    ap.add_argument("--n", type=int, default=6)
    ap.add_argument("--rounds", type=int, default=2)
    ap.add_argument("--temperature", type=float, default=0.95)
    ap.add_argument("--target", type=float, default=0.35, help="AI 分低于它就不再重改")
    ap.add_argument("--no-judge", action="store_true", help="不加载裁判团，直接取第一个候选")
    args = ap.parse_args()

    src = Path(args.input)
    dst = Path(args.output or src.with_name(src.stem + ".orai" + src.suffix))
    paras = read_paragraphs(src)
    todo = [i for i, p in enumerate(paras) if len(p.strip()) >= MIN_LEN]
    current = {i: paras[i].strip() for i in todo}
    report = {i: {"orig": paras[i]} for i in todo}

    for rnd in range(args.rounds if not args.no_judge else 1):
        if not todo:
            break
        from orai.infer.generate import Rewriter
        rw = Rewriter(args.adapter)
        cands = rw.generate([current[i] for i in todo], style=args.style, n=1 if args.no_judge else args.n,
                            temperature=args.temperature)
        del rw
        free_gpu()
        if args.no_judge:
            for i, c in zip(todo, cands):
                current[i] = c[0]
            break
        from orai.infer.judge import Judge
        judge = Judge(with_classifier=Path("models/orai-cls").exists())
        nxt = []
        for i, c in zip(todo, cands):
            # 语义门槛始终对照原段落，避免多轮改写累积跑偏
            ranked = judge.rank(paras[i].strip(), c + ([current[i]] if rnd > 0 else []))
            best = ranked[0]
            if best["ok"]:
                current[i] = best["text"]
            report[i].update({"round": rnd + 1, "out": current[i], "ai_score": best["ai_score"],
                              "ok": best["ok"], "sim": best["sem"]["sim"]})
            if best["ai_score"] > args.target:
                nxt.append(i)
        del judge
        free_gpu()
        print(f"第 {rnd + 1} 轮：{len(todo)} 段，仍偏 AI 的 {len(nxt)} 段", flush=True)
        todo = nxt

    out = list(paras)
    for i, t in current.items():
        out[i] = t
    write_output(src, dst, out)
    with open(dst.with_suffix(".report.jsonl"), "w", encoding="utf-8") as f:
        for i in sorted(report):
            f.write(json.dumps({"para": i, **report[i]}, ensure_ascii=False) + "\n")
    print("输出：", dst)


if __name__ == "__main__":
    main()
